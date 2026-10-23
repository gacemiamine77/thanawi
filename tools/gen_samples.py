#!/usr/bin/env python3
"""Génère de vrais fichiers d'exemple téléchargeables (PDF, DOCX, EPUB)
dans le dossier files/ pour démontrer la fonction téléchargement du site.
Réexécutable : python3 tools/gen_samples.py
"""
import os
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FILES = os.path.join(ROOT, "files")
os.makedirs(FILES, exist_ok=True)

# ---------------------------------------------------------------- PDF -----
from fpdf import FPDF


def make_pdf(path, title, subtitle, sections):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_title(title)
    pdf.set_author("Thanawi - Plateforme du secondaire")

    def mc(h, txt, **kw):
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, h, txt, new_x="LMARGIN", new_y="NEXT", **kw)

    pdf.add_page()
    pdf.set_fill_color(11, 107, 58)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("helvetica", "B", 17)
    mc(10, "  " + title, fill=True)
    pdf.set_text_color(90, 90, 90)
    pdf.set_font("helvetica", "I", 10.5)
    mc(6, subtitle)
    pdf.ln(3)
    pdf.set_text_color(20, 35, 27)
    for h, paras in sections:
        pdf.set_font("helvetica", "B", 13)
        pdf.set_text_color(11, 107, 58)
        mc(8, h)
        pdf.set_text_color(20, 35, 27)
        pdf.set_font("helvetica", "", 11)
        for p in paras:
            mc(6, p)
            pdf.ln(1.5)
        pdf.ln(2)
    pdf.set_font("helvetica", "I", 9)
    pdf.set_text_color(120, 120, 120)
    mc(5, "Document genere par la plateforme Thanawi - usage pedagogique libre.")
    pdf.output(path)
    print("PDF :", os.path.basename(path), os.path.getsize(path), "octets")


make_pdf(
    os.path.join(FILES, "cours-exp-ln.pdf"),
    "Fonctions exponentielle et logarithme neperien",
    "Cours complet - Terminale (3AS) - Mathematiques - Thanawi",
    [
        ("1. La fonction exponentielle", [
            "La fonction exponentielle, notee exp(x) = e^x, est definie sur R, strictement positive et strictement croissante. Elle verifie : e^0 = 1 et (e^x)' = e^x.",
            "Proprietes algebriques : e^(a+b) = e^a x e^b ;  e^(a-b) = e^a / e^b ;  (e^a)^n = e^(na) ;  e^(-x) = 1 / e^x.",
            "Limites : quand x tend vers +infini, e^x tend vers +infini ; quand x tend vers -infini, e^x tend vers 0. Croissances comparees : e^x / x^n tend vers +infini en +infini, et x^n . e^x tend vers 0 en -infini.",
        ]),
        ("2. La fonction logarithme neperien", [
            "La fonction logarithme neperien, notee ln, est definie sur ]0 ; +infini[, strictement croissante, avec ln(1) = 0 et (ln x)' = 1/x. C'est la reciproque de la fonction exponentielle : ln(e^x) = x et, pour x > 0, e^(ln x) = x.",
            "Proprietes (a, b > 0) : ln(ab) = ln a + ln b ;  ln(a/b) = ln a - ln b ;  ln(a^n) = n ln a ;  ln(racine(a)) = (1/2) ln a.",
            "Limites : ln x tend vers +infini en +infini et vers -infini en 0+. Croissances comparees : (ln x)/x tend vers 0 en +infini et x . ln x tend vers 0 en 0+.",
        ]),
        ("3. Equations et inequations (methode)", [
            "e^x = a (a > 0) equivaut a x = ln a.  ln x = a equivaut a x = e^a (avec x > 0).",
            "Les deux fonctions etant strictement croissantes : e^x < e^y equivaut a x < y, et ln x < ln y equivaut a 0 < x < y. Ne jamais oublier les conditions d'existence pour le logarithme.",
        ]),
        ("4. Exemples corriges", [
            "Exemple 1 : resoudre e^(2x) - 3e^x + 2 = 0. On pose X = e^x : X^2 - 3X + 2 = 0, donc X = 1 ou X = 2. D'ou x = 0 ou x = ln 2.",
            "Exemple 2 : resoudre ln(x - 1) = 2. Condition : x > 1. Alors x - 1 = e^2, soit x = e^2 + 1, qui verifie la condition.",
            "Exemple 3 : limite de (e^x - 1)/x en 0. C'est le taux d'accroissement de exp en 0, donc la limite vaut exp'(0) = 1.",
        ]),
        ("5. Exercices d'application", [
            "1) Resoudre dans R : e^(x+1) = 5 ;  ln(2x - 3) = 0 ;  e^(2x) - e^x - 6 = 0.",
            "2) Etudier les variations de f(x) = x - ln x sur ]0 ; +infini[ et en deduire que ln x <= x - 1.",
            "3) Calculer les limites : (ln x)/(x^2) en +infini ;  x^3 . e^x en -infini ;  (e^(2x) - 1)/x en 0.",
            "Les corrections detaillees de ces exercices sont disponibles sur la plateforme Thanawi (bouton Solution).",
        ]),
    ],
)

make_pdf(
    os.path.join(FILES, "sujet-rc.pdf"),
    "Sujet type Bac - Physique : circuit RC",
    "Terminale (3AS) - Sciences experimentales - Duree conseillee : 1 h - Sujet d'entrainement redige pour Thanawi (ce n'est pas un sujet officiel)",
    [
        ("Exercice : charge d'un condensateur (7 points)", [
            "On considere un circuit serie comportant un generateur ideal de tension E = 6 V, un conducteur ohmique de resistance R = 1,0 kohm, un condensateur de capacite C inconnue et un interrupteur K. A l'instant t = 0, on ferme l'interrupteur, le condensateur etant initialement decharge.",
            "1) Faire le schema du circuit et flécher la tension uC aux bornes du condensateur et l'intensite i du courant.",
            "2) Etablir l'equation differentielle verifiee par uC(t) pendant la charge.",
            "3) Verifier que uC(t) = E (1 - e^(-t/RC)) est solution de cette equation differentielle.",
            "4) Definir la constante de temps tau du circuit et donner son expression. Par analyse dimensionnelle, verifier que tau a la dimension d'un temps.",
            "5) L'experience donne uC = 3,8 V a l'instant t = 10 ms. Determiner la valeur de tau, puis celle de la capacite C. (On rappelle que pour t = tau, uC atteint environ 63 % de sa valeur finale.)",
            "6) Au bout de quelle duree peut-on considerer le condensateur comme totalement charge ? Quelle est alors l'energie emmagasinee dans le condensateur ?",
        ]),
        ("Donnees utiles", [
            "Relation charge-tension : q = C . uC.  Intensite : i = dq/dt.  Energie d'un condensateur : Eelec = (1/2) C uC^2.  1 kohm = 10^3 ohm.",
            "Le corrige type est disponible sur la plateforme Thanawi (fichier separe a telecharger apres avoir cherche le sujet).",
        ]),
    ],
)

make_pdf(
    os.path.join(FILES, "sujet-rc-corrige.pdf"),
    "Corrige type - Physique : circuit RC",
    "Correction du sujet d'entrainement Thanawi - Terminale (3AS)",
    [
        ("Correction detaillee", [
            "2) Loi des mailles pendant la charge : E = uR + uC = R . i + uC. Or i = dq/dt et q = C uC, donc i = C duC/dt. D'ou l'equation differentielle : RC duC/dt + uC = E.",
            "3) Si uC = E(1 - e^(-t/RC)), alors duC/dt = (E/RC) e^(-t/RC). Donc RC duC/dt + uC = E e^(-t/RC) + E - E e^(-t/RC) = E : la fonction est bien solution. De plus uC(0) = 0, conforme a la condition initiale.",
            "4) La constante de temps est tau = RC. Dimension : [R] = [U]/[I] et [C] = [Q]/[U] = [I][T]/[U], donc [RC] = [T] : tau est bien un temps. Graphiquement, tau est l'abscisse du point d'intersection de la tangente a l'origine avec l'asymptote uC = E.",
            "5) 3,8 / 6 = 0,63 : a t = 10 ms, uC vaut 63 % de E, donc tau = 10 ms = 0,010 s. Alors C = tau / R = 0,010 / 1000 = 1,0 x 10^-5 F = 10 microF.",
            "6) Le condensateur est pratiquement charge au bout de t = 5 tau = 50 ms (uC atteint alors plus de 99 % de E). Energie finale : Eelec = (1/2) C E^2 = 0,5 x 1,0x10^-5 x 36 = 1,8 x 10^-4 J.",
        ]),
        ("Conseils de methode", [
            "Toujours commencer par la loi des mailles et la convention choisie (schema fléche). Verifier les conditions initiales quand on propose une solution. Connaitre les valeurs remarquables : a t = tau, 63 % de la charge ; a t = 5 tau, regime permanent pratiquement atteint.",
        ]),
    ],
)

# --------------------------------------------------------------- DOCX -----
from docx import Document
from docx.shared import Pt, RGBColor

doc = Document()
t = doc.add_heading("Série d'exercices — Équations du second degré", level=1)
for r in t.runs:
    r.font.color.rgb = RGBColor(0x0B, 0x6B, 0x3A)
p = doc.add_paragraph()
p.add_run("2e année secondaire (2AS) — Mathématiques — Plateforme Thanawi\n").italic = True
p.add_run("Rappel : pour ax² + bx + c = 0 (a ≠ 0), le discriminant est Δ = b² − 4ac. "
          "Si Δ > 0 : deux solutions x = (−b ± √Δ) / 2a. Si Δ = 0 : une solution double x = −b / 2a. "
          "Si Δ < 0 : pas de solution réelle. Somme S = −b/a, produit P = c/a.")

exercices = [
    ("Exercice 1", "Résoudre dans ℝ : a) x² − 5x + 6 = 0 ;  b) 2x² + 3x − 2 = 0 ;  c) x² + 2x + 5 = 0 ;  d) 4x² − 12x + 9 = 0."),
    ("Exercice 2", "Sans calculer le discriminant, trouver deux nombres de somme 7 et de produit 12. Même question avec S = −3 et P = −10."),
    ("Exercice 3", "Étudier le signe de f(x) = −x² + 4x − 3, puis résoudre l'inéquation f(x) ≥ 0."),
    ("Exercice 4", "Un rectangle a pour périmètre 34 m et pour aire 60 m². Déterminer ses dimensions."),
    ("Exercice 5", "Pour quelles valeurs du paramètre m l'équation x² − 2(m + 1)x + m² = 0 admet-elle une solution double ? Deux solutions distinctes ?"),
]
for h, txt in exercices:
    doc.add_heading(h, level=2)
    doc.add_paragraph(txt)

p = doc.add_paragraph()
r = p.add_run("Les solutions détaillées sont sur la plateforme Thanawi : ouvrez cette ressource et cliquez sur « Afficher la solution » après avoir cherché.")
r.italic = True
r.font.size = Pt(10)
docx_path = os.path.join(FILES, "exercices-second-degre.docx")
doc.save(docx_path)
print("DOCX :", os.path.basename(docx_path), os.path.getsize(docx_path), "octets")

# --------------------------------------------------------------- EPUB -----
epub_path = os.path.join(FILES, "methodo-dissertation.epub")
container = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>"""
opf = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" unique-identifier="bid" version="3.0">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="bid">urn:uuid:thanawi-methodo-dissertation</dc:identifier>
    <dc:title>Méthodologie de la dissertation philosophique</dc:title>
    <dc:creator>Thanawi</dc:creator><dc:language>fr</dc:language>
    <meta property="dcterms:modified">2026-10-23T00:00:00Z</meta>
  </metadata>
  <manifest>
    <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
    <item id="c1" href="chap1.xhtml" media-type="application/xhtml+xml"/>
    <item id="c2" href="chap2.xhtml" media-type="application/xhtml+xml"/>
    <item id="c3" href="chap3.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="nav"/><itemref idref="c1"/><itemref idref="c2"/><itemref idref="c3"/></spine>
</package>"""
nav = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="fr" lang="fr">
<head><title>Sommaire</title></head><body>
<h1>Méthodologie de la dissertation philosophique</h1>
<p>Résumé méthodologique — Terminale (3AS), filière Lettres et philosophie — Plateforme Thanawi.</p>
<nav epub:type="toc"><ol>
<li><a href="chap1.xhtml">1. Analyser le sujet et construire la problématique</a></li>
<li><a href="chap2.xhtml">2. Les plans : dialectique, analytique, comparatif</a></li>
<li><a href="chap3.xhtml">3. Rédiger : introduction, développement, conclusion</a></li>
</ol></nav></body></html>"""


def chap(title, paras):
    body = "".join(f"<p>{p}</p>" for p in paras)
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xml:lang="fr" lang="fr">
<head><title>{title}</title></head><body><h2>{title}</h2>{body}</body></html>"""


chaps = [
    ("chap1.xhtml", chap("1. Analyser le sujet et construire la problématique", [
        "Commencez par définir les termes clés du sujet : la plupart des erreurs viennent d'une lecture trop rapide. Repérez la forme du sujet (question, affirmation à discuter, comparaison).",
        "La problématique est la tension cachée du sujet : deux réponses opposées semblent possibles. Formulez-la en une phrase interrogative qui montre le conflit, par exemple : « La conscience est-elle un accès certain à soi-même, ou peut-elle nous tromper ? »",
        "Astuce : si vous pouvez répondre immédiatement « oui » ou « non » sans hésiter, votre problématique est probablement trop faible. Cherchez ce qui rend la question réellement difficile.",
    ])),
    ("chap2.xhtml", chap("2. Les plans : dialectique, analytique, comparatif", [
        "Plan dialectique (thèse – antithèse – synthèse) : adapté aux sujets qui appellent une discussion. On expose une position, puis son contraire, puis on dépasse l'opposition par une position plus nuancée.",
        "Plan analytique : on décompose la notion en ses aspects (causes, aspects, conséquences). Utile pour les sujets-notions : « Qu'est-ce que la justice ? »",
        "Plan comparatif : pour les sujets qui comparent deux domaines (science et philosophie, travail et aliénation…). On étudie les ressemblances, puis les différences, puis le rapport entre les deux termes.",
        "Dans tous les cas : chaque partie doit contenir des arguments précis, un ou deux exemples, et si possible une référence philosophique bien comprise (mieux vaut un auteur maîtrisé que cinq noms cités sans compréhension).",
    ])),
    ("chap3.xhtml", chap("3. Rédiger : introduction, développement, conclusion", [
        "Introduction (4 temps) : accroche générale sur le thème, présentation et analyse du sujet, problématique, annonce du plan.",
        "Développement : une idée par paragraphe, des transitions qui montrent la progression (« Cependant… », « Mais cette position rencontre une limite… »). Évitez de juxtaposer les opinions sans les affronter.",
        "Conclusion : répondez clairement à la problématique, résumez le chemin parcouru, puis ouvrez sur une question connexe. N'introduisez aucun argument nouveau.",
        "Critères d'évaluation au baccalauréat : compréhension du sujet, qualité de la problématique, rigueur argumentative, culture philosophique et clarté de l'expression. Relisez toujours votre copie en gardant 10 minutes pour cela.",
    ])),
]

with zipfile.ZipFile(epub_path, "w") as z:
    z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
    z.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
    z.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
    z.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
    for name, content in chaps:
        z.writestr("OEBPS/" + name, content, compress_type=zipfile.ZIP_DEFLATED)
print("EPUB :", os.path.basename(epub_path), os.path.getsize(epub_path), "octets")
print("Terminé. Fichiers dans :", FILES)
